#!/usr/bin/python3
"""The gateway's .xlsx writer and the Excel stand-in's staging format.
    python3 tests/xlsx_test.py        (needs python3-openpyxl)"""
import importlib.machinery
import importlib.util
import os
import sys
import tempfile
import types
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
for mod in ("websockify", "websockify.websocketproxy"):
    sys.modules[mod] = types.ModuleType(mod)
sys.modules["websockify.websocketproxy"].ProxyRequestHandler = object
sys.modules["websockify.websocketproxy"].WebSocketProxy = object
loader = importlib.machinery.SourceFileLoader(
    "gateway", os.path.join(HERE, "..", "rootfs-vnc", "usr", "local", "bin", "wine-webapp-gateway"))
spec = importlib.util.spec_from_loader("gateway", loader)
gw = importlib.util.module_from_spec(spec)
loader.exec_module(gw)
import openpyxl  # noqa: E402


class Xlsx(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()

    def test_cells_types_and_formats(self):
        cells = [(1, 1, "s", "", "Report"), (3, 1, "s", "", "Amount"),
                 (4, 1, "n", "#,##0.00", "1234.5"), (5, 1, "n", "", "7"),
                 (6, 1, "s", "", 'two\nlines <&> "q"'), (7, 1, "b", "", "True"),
                 (8, 2, "n", "", "not-a-number")]
        path = os.path.join(self.d, "t.xlsx")
        gw.write_xlsx(cells, path)
        ws = openpyxl.load_workbook(path).active
        self.assertEqual(ws["A1"].value, "Report")
        self.assertEqual(ws["A4"].value, 1234.5)
        self.assertEqual(ws["A4"].number_format, "#,##0.00")
        self.assertEqual(ws["A5"].value, 7)
        self.assertEqual(ws["A6"].value, 'two\nlines <&> "q"')
        self.assertIs(ws["A7"].value, True)
        self.assertEqual(ws["B8"].value, "not-a-number")

    def test_staging_file_becomes_a_listed_spreadsheet(self):
        exports, prints = os.path.join(self.d, "x"), os.path.join(self.d, "p")
        os.makedirs(exports)
        with open(os.path.join(exports, "1-2.cells"), "w") as f:
            f.write("1\t1\ts\t\tOpen Listing: 2026\n2\t1\tn\t0.00\t3.5\n2\t2\ts\t\ta\\tb\\nc\\\\d\n")
        gw.EXPORT_DIR, gw.PRINT_DIR = exports, prints
        gw.convert_exports()
        names = os.listdir(prints)
        self.assertEqual(len(names), 1)
        self.assertRegex(names[0], r"^Open_Listing_2026-xlsx_\d+\.xlsx$")
        self.assertEqual(os.listdir(exports), [])
        ws = openpyxl.load_workbook(os.path.join(prints, names[0])).active
        self.assertEqual(ws["B2"].value, "a\tb\nc\\d")
        self.assertEqual(ws["A2"].value, 3.5)


if __name__ == "__main__":
    unittest.main(verbosity=1)
